"""Build sentences.jsonl from Shah et al. FOMC hawkish/dovish annotations.

Label mapping (from gtfintechlab README / HuggingFace FOMC-RoBERTa):
  0 = dovish, 1 = hawkish, 2 = neutral
Uses annotated split workbooks under data/raw/shah/data/annotated_data/.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from strip_meta import strip_meta  # noqa: E402

ANNOTATED = ROOT / "data" / "raw" / "shah" / "data" / "annotated_data"
OUT = ROOT / "data" / "clean" / "sentences.jsonl"

LABEL_MAP = {0: "dovish", 1: "hawkish", 2: "neutral", "0": "dovish", "1": "hawkish", "2": "neutral"}

FILES = [
    ("manual-sp-split.xlsx", "speech"),
    ("manual-mm-split.xlsx", "minutes"),
    ("manual-pc-split.xlsx", "press_conference"),
]


def load_rows(path: Path) -> list[tuple]:
    wb = load_workbook(path, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h).lower() if h is not None else "" for h in rows[0]]
    # columns: index?, sentence, year, label, orig_index?
    sent_i = hdr.index("sentence") if "sentence" in hdr else 1
    year_i = hdr.index("year") if "year" in hdr else 2
    label_i = hdr.index("label") if "label" in hdr else 3
    out = []
    for r in rows[1:]:
        if not r or len(r) <= label_i:
            continue
        sent = r[sent_i]
        lab = r[label_i]
        year = r[year_i]
        if sent is None or lab in (None, "-", ""):
            continue
        if lab not in LABEL_MAP:
            continue
        out.append((str(sent).strip(), LABEL_MAP[lab], year))
    return out


def main() -> None:
    records = []
    seen = set()
    for fname, doc_type in FILES:
        path = ANNOTATED / fname
        rows = load_rows(path)
        for i, (sent, label, year) in enumerate(rows):
            key = (sent.lower(), label)
            if key in seen:
                continue
            seen.add(key)
            sent_id = f"shah-{doc_type}-{year}-{i:04d}"
            records.append(
                {
                    "sent_id": sent_id,
                    "text": strip_meta(sent),
                    "label": label,
                    "raw_text": sent,
                    "year": int(year) if year is not None else None,
                    "doc_type": doc_type,
                    "source": "shah",
                }
            )
        print(f"{fname}: kept unique from {len(rows)} rows; cumulative {len(records)}")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w") as f:
        for r in records:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")

    from collections import Counter
    print("label counts", Counter(r["label"] for r in records))
    print(f"Wrote {len(records)} -> {OUT}")


if __name__ == "__main__":
    main()
